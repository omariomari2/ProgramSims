
import openpyxl
import os

filepath = "Delinquency_prediction_dataset.xlsx"

def inspect_workbook():
    try:
        wb = openpyxl.load_workbook(filepath, read_only=True)
        print(f"Sheet names (read_only=True): {wb.sheetnames}")
    except Exception as e:
        print(f"Error (read_only=True): {e}")

    try:
        wb = openpyxl.load_workbook(filepath, read_only=False)
        print(f"Sheet names (read_only=False): {wb.sheetnames}")
    except Exception as e:
        print(f"Error (read_only=False): {e}")

if __name__ == "__main__":
    if os.path.exists(filepath):
        inspect_workbook()
